import copy, os, tempfile, unittest, zipfile
from pathlib import Path
from docx import Document as PyDoc

from explain_core import adapt_text, adapt_text_for_target, analyze, export, unpack_targets, zip_output
from workbench.docxio import Package, resolve_revisions, w, text, local

class ExplainCoreTests(unittest.TestCase):
    def test_voice_rules_are_narrow(self):
        self.assertEqual(adapt_text('截至募集说明书签署日，发行人无重大变化。'),'截至本说明性文件出具日，公司无重大变化。')
        self.assertEqual(adapt_text('具体情况详见募集说明书“第五节”。'),'具体情况详见募集说明书“第五节”。')
        self.assertEqual(adapt_text('本募集说明书所称发行人'),'募集说明书所称公司')
        self.assertEqual(adapt_text('截至说明出具日，公司情况如下'),'截至本说明性文件出具日，公司情况如下')
        self.assertEqual(adapt_text('截至本说明出具日，公司情况如下'),'截至本说明性文件出具日，公司情况如下')
        self.assertEqual(adapt_text_for_target('（十六）2023年1月，发行人发生事项。','2023年1月，公司发生事项。'),'2023年1月，公司发生事项。')

    def make_doc(self,path,paras=(),table=None):
        d=PyDoc()
        for p in paras:d.add_paragraph(p)
        if table:
            t=d.add_table(rows=len(table),cols=len(table[0]))
            for i,row in enumerate(table):
                for j,v in enumerate(row):t.cell(i,j).text=str(v)
        d.save(path)


    def test_zip_output_preserves_real_mtime_and_utf8_names(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);folder=td/'out';folder.mkdir();fp=folder/'中文说明.docx';fp.write_bytes(b'test')
            os.utime(fp,(1760000000,1760000000))
            zp=td/'x.zip';zip_output(folder,zp)
            with zipfile.ZipFile(zp) as z:
                info=z.getinfo('中文说明.docx')
                self.assertNotEqual(info.date_time[:3],(1980,1,1))
                self.assertTrue(info.flag_bits & 0x800)

    def test_source_and_wording_proposals(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx'
            self.make_doc(src,['河南测试有限公司2026年公司债券募集说明书','第四节 发行人基本情况','截至募集说明书签署日，发行人注册资本为100亿元。'])
            self.make_doc(tar,['发行人基本情况说明','河南测试有限公司（以下简称“公司”）说明如下：','截至本说明性文件出具日，公司注册资本为90亿元。','特此说明。','（本页无正文，为《发行人基本情况说明》之盖章页）'])
            r=analyze(src,[(str(tar),tar.name)])
            self.assertTrue(any(p['status']=='ready' and '90亿元' in p['old_text'] for p in r['proposals']))
            # Main document title and seal-page title reference are frozen identity,
            # not body wording; they must not become wording proposals.
            self.assertFalse(any('发行人基本情况说明' in p['old_text'] for p in r['proposals']))

    def test_main_title_and_seal_page_are_frozen_from_wording_sweep(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            title='发行人关于主要业务板块行业地位及竞争情况的说明文件'
            seal=f'（本页无正文，为《{title}》之盖章页）'
            self.make_doc(src,['河南测试有限公司募集说明书','第四节 主营业务','发行人主营业务保持稳定。'])
            self.make_doc(tar,[title,'河南测试有限公司（以下简称“公司”）说明如下：','发行人主营业务保持稳定。','特此说明。',seal])
            r=analyze(src,[(str(tar),tar.name)])
            self.assertFalse(any(p['old_text']==title for p in r['proposals']))
            self.assertFalse(any(p['old_text']==seal for p in r['proposals']))
            body=next(p for p in r['proposals'] if '主营业务保持稳定' in p['old_text'])
            self.assertEqual(body['status'],'ready')
            rep=export(r,{},out,'柒');self.assertFalse(rep['blocked']);self.assertFalse(rep['errors'])
            pkg=Package(out/tar.name);accept=copy.deepcopy(pkg.root);resolve_revisions(accept,True)
            vis='\n'.join(text(e) for e in accept.find(w('body')))
            self.assertIn(title,vis)
            self.assertIn(seal,vis)
            self.assertIn('公司主营业务保持稳定。',vis)
            self.assertNotIn('公司关于主要业务板块行业地位及竞争情况的说明文件',vis)

    def test_export_tracks_whole_block_and_reject_restores(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            self.make_doc(src,['河南测试有限公司2026年公司债券募集说明书','第四节 发行人基本情况','截至募集说明书签署日，发行人注册资本为100亿元。'])
            self.make_doc(tar,['公司基本情况说明','河南测试有限公司（以下简称“公司”）说明如下：','截至本说明性文件出具日，公司注册资本为90亿元。','特此说明。'])
            baseline=Package(tar).baseline()
            def sem(root):
                body=root.find(w('body'))
                return [(local(e),text(e).strip()) for e in body if e.tag in (w('p'),w('tbl')) and text(e).strip()]
            base_sem=sem(copy.deepcopy(baseline.root))
            r=analyze(src,[(str(tar),tar.name)]);rep=export(r,{},out,'柒');self.assertFalse(rep['errors'])
            fp=out/tar.name;pkg=Package(fp);reject=copy.deepcopy(pkg.root);resolve_revisions(reject,False)
            self.assertEqual(sem(reject),base_sem)
            accept=copy.deepcopy(pkg.root);resolve_revisions(accept,True)
            vis='\n'.join(text(e) for e in accept.find(w('body')))
            self.assertIn('100亿元',vis);self.assertNotIn('90亿元',vis);self.assertIn('本说明性文件出具日',vis);self.assertNotIn('发行人注册资本',vis)

    def test_numbered_fact_heading_refreshes_but_keeps_target_prefix(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            ds=PyDoc();ds.add_paragraph('河南测试有限公司募集说明书');ds.add_heading('（三）债券情况',level=2);ds.add_heading('2、截至募集说明书签署日，发行人已发行债券258只。',level=3);ds.save(src)
            dt=PyDoc();dt.add_paragraph('债券情况说明');dt.add_heading('（三）公司债券情况',level=2);dt.add_heading('2、截至说明出具日，公司已发行债券256只。',level=3);dt.add_paragraph('特此说明。');dt.save(tar)
            r=analyze(src,[(str(tar),tar.name)])
            hits=[p for p in r['proposals'] if '256只' in p['old_text']]
            self.assertEqual(len(hits),1);self.assertEqual(hits[0]['status'],'ready')
            export(r,{},out,'柒');pkg=Package(out/tar.name);accept=copy.deepcopy(pkg.root);resolve_revisions(accept,True)
            vis='\n'.join(text(e) for e in accept.find(w('body')))
            self.assertIn('2、截至本说明性文件出具日，公司已发行债券258只。',vis)

    def test_short_numbered_list_item_is_source_fact_not_structural_heading(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            self.make_doc(src,['河南测试有限公司募集说明书','1、股东会','股东会行使以下职权：','（4）审议批准公司的年度财务预算方案、决算方案；'])
            self.make_doc(tar,['治理结构说明','河南测试有限公司（以下简称“公司”）说明如下：','1、股东会','股东会行使以下职权：','（4）审批批准公司的年度财务预算方案、决算方案；','特此说明。'])
            r=analyze(src,[(str(tar),tar.name)])
            hit=[p for p in r['proposals'] if '审批批准' in p['old_text']]
            self.assertEqual(len(hit),1)
            self.assertEqual(hit[0]['status'],'ready')
            rep=export(r,{},out,'柒');self.assertFalse(rep['blocked']);self.assertFalse(rep['errors'])
            pkg=Package(out/tar.name);accept=copy.deepcopy(pkg.root);resolve_revisions(accept,True)
            vis='\n'.join(text(e) for e in accept.find(w('body')))
            self.assertIn('（4）审议批准公司的年度财务预算方案、决算方案；',vis)
            self.assertNotIn('审批批准',vis)

    def test_source_local_number_is_not_imported_into_unnumbered_target(self):
        self.assertEqual(adapt_text_for_target('（十六）2023年1月，发行人召开会议。','2023年1月，公司召开会议。'),'2023年1月，公司召开会议。')

    def test_table_is_single_object_proposal(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx'
            self.make_doc(src,['河南测试有限公司2026年公司债券募集说明书','授信情况如下：'],[['银行','额度'],['甲行','100']])
            self.make_doc(tar,['授信说明','河南测试有限公司（以下简称“公司”）说明如下：','授信情况如下：'],[['银行','额度'],['甲行','80']])
            d=PyDoc(tar);d.add_paragraph('特此说明。');d.save(tar)
            r=analyze(src,[(str(tar),tar.name)])
            tabs=[p for p in r['proposals'] if p['kind']=='table']
            self.assertEqual(len(tabs),1);self.assertEqual(tabs[0]['status'],'ready')


    def test_identical_source_owned_still_rematerializes(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            self.make_doc(src,['河南测试有限公司募集说明书','第四节 基本情况','截至募集说明书签署日，发行人注册资本为100亿元。'])
            self.make_doc(tar,['基本情况说明','河南测试有限公司（以下简称“公司”）说明如下：','截至本说明性文件出具日，公司注册资本为100亿元。','特此说明。'])
            r=analyze(src,[(str(tar),tar.name)])
            p=next(p for p in r['proposals'] if '注册资本为100亿元' in p['old_text'])
            self.assertEqual(p['status'],'ready');self.assertTrue(p['same_visible'])
            rep=export(r,{},out,'柒')
            rec=next(a for a in rep['files'][0]['applied_records'] if a['id']==p['id'])
            self.assertEqual(rec['mode'],'source_confirmed_refresh')
            pkg=Package(out/tar.name);accept=copy.deepcopy(pkg.root);resolve_revisions(accept,True)
            reject=copy.deepcopy(pkg.root);resolve_revisions(reject,False)
            av='\n'.join(text(e) for e in accept.find(w('body')))
            rv='\n'.join(text(e) for e in reject.find(w('body')))
            self.assertIn('截至本说明性文件出具日，公司注册资本为100亿元。',av)
            self.assertIn('截至本说明性文件出具日，公司注册资本为100亿元。',rv)

    def test_identical_source_owned_uses_source_object_not_target_clone(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            ds=PyDoc();ds.add_paragraph('河南测试有限公司募集说明书');ds.add_paragraph('第四节 基本情况')
            ps=ds.add_paragraph();rs=ps.add_run('截至募集说明书签署日，发行人注册资本为100亿元。');rs.bold=True
            ds.save(src)
            dt=PyDoc();dt.add_paragraph('基本情况说明');dt.add_paragraph('河南测试有限公司（以下简称“公司”）说明如下：')
            pt=dt.add_paragraph();rt=pt.add_run('截至本说明性文件出具日，公司注册资本为100亿元。');rt.italic=True
            dt.add_paragraph('特此说明。');dt.save(tar)
            r=analyze(src,[(str(tar),tar.name)])
            p=next(p for p in r['proposals'] if '注册资本为100亿元' in p['old_text'])
            self.assertTrue(p['same_visible'])
            rep=export(r,{},out,'柒');self.assertFalse(rep['blocked']);self.assertFalse(rep['errors'])
            pkg=Package(out/tar.name);accept=copy.deepcopy(pkg.root);resolve_revisions(accept,True)
            target_para=next(e for e in accept.find(w('body')) if e.tag==w('p') and '注册资本为100亿元' in text(e))
            runs=target_para.findall('.//'+w('r'))
            self.assertTrue(any(r.find(w('rPr')+'/'+w('b')) is not None for r in runs), 'same-visible refresh must carry source run formatting')
            self.assertFalse(any(r.find(w('rPr')+'/'+w('i')) is not None for r in runs), 'target run formatting must not be cloned as source content')

    def test_target_schema_table_projection_is_whole_object(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            self.make_doc(src,['河南测试有限公司募集说明书','募集资金用途如下：'],[
                ['发行主体','证券简称','债券类型','起息日','到期日','发行规模'],
                ['发行人','24测试01','公募公司债','2024-1-1','2027-1-1','10.00'],
                ['合计','','','','','10.00']])
            self.make_doc(tar,['募集资金说明','河南测试有限公司（以下简称“公司”）说明如下：','募集资金用途如下：'],[
                ['证券简称','起息日','到期日','发行规模'],
                ['24测试01','2024-1-1','2027-1-1','9.00'],
                ['合计','','','9.00']])
            d=PyDoc(tar);d.add_paragraph('特此说明。');d.save(tar)
            r=analyze(src,[(str(tar),tar.name)])
            p=next(p for p in r['proposals'] if p['kind']=='table')
            self.assertEqual(p['status'],'ready');self.assertEqual(p['table_projection'],[1,3,4,5])
            rep=export(r,{},out,'柒');self.assertFalse(rep['blocked']);self.assertFalse(rep['errors'])
            rec=next(a for a in rep['files'][0]['applied_records'] if a['id']==p['id'])
            self.assertEqual(rec['mode'],'source_table_projection')
            pkg=Package(out/tar.name);accept=copy.deepcopy(pkg.root);resolve_revisions(accept,True)
            tables=accept.findall('.//'+w('tbl'))
            projected=None
            for tbl in tables:
                rows=[[text(tc).strip() for tc in tr.findall(w('tc'))] for tr in tbl.findall(w('tr'))]
                if rows and rows[0] and rows[0][0]=='证券简称':projected=rows;break
            self.assertIsNotNone(projected)
            self.assertEqual(projected[0],['证券简称','起息日','到期日','发行规模'])
            self.assertEqual(projected[1],['24测试01','2024-1-1','2027-1-1','10.00'])

    def test_mixed_analysis_table_is_review_not_whole_copy(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx'
            self.make_doc(src,['河南测试有限公司募集说明书','毛利率情况如下：'],[
                ['项目','2025年度','2024年度'],['业务A','10%','9%'],['合计','10%','9%']])
            self.make_doc(tar,['毛利率说明','河南测试有限公司（以下简称“公司”）说明如下：','毛利率情况如下：'],[
                ['项目','2025年度','2024年度'],['业务A','10%','9%'],['同行业对比情况','8%','7%']])
            d=PyDoc(tar);d.add_paragraph('特此说明。');d.save(tar)
            r=analyze(src,[(str(tar),tar.name)])
            p=next(p for p in r['proposals'] if p['kind']=='table')
            self.assertEqual(p['status'],'retained');self.assertEqual(p['decision'],'keep');self.assertFalse(p['review_required'])

    def test_projection_can_refresh_a_changed_percentage_from_source_clause(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            self.make_doc(src,['河南测试有限公司募集说明书',
                '2024年3月，根据工作部署，公司发生股权划转。河南省财政厅享有70.000%表决权，为公司控股股东及实际控制人。'])
            self.make_doc(tar,['股权情况说明','河南测试有限公司（以下简称“公司”）说明如下：',
                '截至本说明性文件出具日，河南省财政厅享有公司68.336%表决权，为公司控股股东及实际控制人。','特此说明。'])
            r=analyze(src,[(str(tar),tar.name)])
            p=next(p for p in r['proposals'] if '68.336%' in p['old_text'])
            self.assertEqual(p['status'],'ready');self.assertTrue(p['projection_mode'])
            self.assertEqual(p.get('projection_source_value'),'70.000%')
            rep=export(r,{},out,'柒');self.assertFalse(rep['blocked']);self.assertFalse(rep['errors'])
            pkg=Package(out/tar.name);accept=copy.deepcopy(pkg.root);resolve_revisions(accept,True)
            vis='\n'.join(text(e) for e in accept.find(w('body')))
            self.assertIn('河南省财政厅享有公司70.000%表决权',vis)
            self.assertNotIn('68.336%',vis)

    def test_unit_line_binds_to_the_selected_source_table(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx'
            ds=PyDoc();ds.add_paragraph('河南测试有限公司募集说明书')
            ds.add_paragraph('A表如下：');ds.add_paragraph('单位：万元')
            t=ds.add_table(rows=2,cols=2);t.cell(0,0).text='项目';t.cell(0,1).text='金额';t.cell(1,0).text='A';t.cell(1,1).text='10'
            ds.add_paragraph('B表如下：');ds.add_paragraph('单位：万元')
            t=ds.add_table(rows=2,cols=2);t.cell(0,0).text='项目';t.cell(0,1).text='金额';t.cell(1,0).text='B';t.cell(1,1).text='20'
            ds.save(src)
            dt=PyDoc();dt.add_paragraph('B表说明');dt.add_paragraph('河南测试有限公司（以下简称“公司”）说明如下：')
            dt.add_paragraph('B表如下：');dt.add_paragraph('单位：万元')
            t=dt.add_table(rows=2,cols=2);t.cell(0,0).text='项目';t.cell(0,1).text='金额';t.cell(1,0).text='B';t.cell(1,1).text='19'
            dt.add_paragraph('特此说明。');dt.save(tar)
            r=analyze(src,[(str(tar),tar.name)])
            unit=next(p for p in r['proposals'] if p['old_text'].strip()=='单位：万元')
            table=next(p for p in r['proposals'] if p['kind']=='table')
            self.assertEqual(table['status'],'ready')
            self.assertEqual(unit['status'],'ready')
            self.assertEqual(unit['selected_source_index'],table['selected_source_index']-1)
            self.assertEqual(len(unit['candidates']),1)

    def test_unit_line_stays_with_review_table_instead_of_global_matching(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx'
            self.make_doc(src,['河南测试有限公司募集说明书','毛利率情况如下：','单位：%'],[
                ['项目','2025年度','2024年度'],['业务A','10%','9%'],['合计','10%','9%']])
            self.make_doc(tar,['毛利率说明','河南测试有限公司（以下简称“公司”）说明如下：','毛利率情况如下：','单位：%'],[
                ['项目','2025年度','2024年度'],['业务A','10%','9%'],['同行业对比情况','8%','7%']])
            d=PyDoc(tar);d.add_paragraph('特此说明。');d.save(tar)
            r=analyze(src,[(str(tar),tar.name)])
            unit=next(p for p in r['proposals'] if p['old_text'].strip()=='单位：%')
            table=next(p for p in r['proposals'] if p['kind']=='table')
            self.assertEqual(table['status'],'retained')
            self.assertEqual(unit['status'],'retained')
            self.assertEqual(unit['decision'],'keep')
            self.assertIsNone(unit['selected_source_index'])

    def test_structural_gap_refreshes_renamed_department_body(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            self.make_doc(src,['河南测试有限公司募集说明书','4、办公室','办公室负责综合协调。','5、战略投资部','战略投资部负责战略研究、投资发展和产业合作。','6、计划运营部','计划运营部负责经营计划。'])
            self.make_doc(tar,['组织机构说明','河南测试有限公司（以下简称“公司”）说明如下：','4、办公室','办公室负责综合协调。','5、战略发展部','战略发展部负责战略规划和产业研究。','6、计划运营部','计划运营部负责经营计划。','特此说明。'])
            r=analyze(src,[(str(tar),tar.name)])
            p=next(p for p in r['proposals'] if '战略发展部负责' in p['old_text'])
            self.assertEqual(p['status'],'ready')
            self.assertTrue(p['source_owned'])
            self.assertIn('结构位置确认',p['reason'])
            rep=export(r,{},out,'柒');self.assertFalse(rep['blocked']);self.assertFalse(rep['errors'])
            pkg=Package(out/tar.name);accept=copy.deepcopy(pkg.root);resolve_revisions(accept,True)
            vis='\n'.join(text(e) for e in accept.find(w('body')))
            self.assertIn('战略投资部负责战略研究、投资发展和产业合作。',vis)
            self.assertNotIn('战略发展部负责战略规划和产业研究。',vis)
            # Target structural headings are not copied/relabelled by the source engine.
            self.assertIn('5、战略发展部',vis)

    def test_clause_projection_uses_exact_source_clause_before_next_anchor(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            self.make_doc(src,['河南测试有限公司募集说明书',
                '截至2026年3月末，甲账户分别持有发行人子公司105,100股和204,200股；乙机构另持有发行人子公司100股和129,000股。',
                '除上述情况外，发行人与有关机构不存在其他重大利害关系。'])
            self.make_doc(tar,['重大利害关系说明','河南测试有限公司（以下简称“公司”）说明如下：',
                '截至2026年3月末，甲账户分别持有公司子公司105,100股和204,200股。',
                '除上述情况外，公司与有关机构不存在其他重大利害关系。','特此说明。'])
            r=analyze(src,[(str(tar),tar.name)])
            p=next(p for p in r['proposals'] if '105,100股' in p['old_text'])
            self.assertEqual(p['status'],'ready');self.assertTrue(p['projection_mode'])
            self.assertIn('募集说明书单一分句',p['reason'])
            rep=export(r,{},out,'柒');self.assertFalse(rep['blocked']);self.assertFalse(rep['errors'])
            pkg=Package(out/tar.name);accept=copy.deepcopy(pkg.root);resolve_revisions(accept,True)
            vis='\n'.join(text(e) for e in accept.find(w('body')))
            self.assertIn('甲账户分别持有公司子公司105,100股和204,200股',vis)
            self.assertNotIn('乙机构另持有公司子公司100股和129,000股',vis)

    def test_clause_projection_does_not_drop_target_only_rationale(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx'
            self.make_doc(src,['河南测试有限公司募集说明书',
                '截至2026年3月末，示例证券持有发行人子公司61,969股；其他机构另有持股。',
                '除上述情况外，发行人与有关机构不存在其他重大利害关系。'])
            self.make_doc(tar,['重大利害关系说明','河南测试有限公司（以下简称“公司”）说明如下：',
                '截至2026年3月末，示例证券持有公司子公司61,969股。该等持股系场外期权交易对冲风险需要，不带有自营择时观点。',
                '除上述情况外，公司与有关机构不存在其他重大利害关系。','特此说明。'])
            r=analyze(src,[(str(tar),tar.name)])
            p=next(p for p in r['proposals'] if '61,969股' in p['old_text'])
            self.assertEqual(p['status'],'retained')
            self.assertEqual(p['decision'],'keep')

    def test_review_update_requires_explicit_current_candidate(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            self.make_doc(src,['河南测试有限公司募集说明书','第四节 情况','截至募集说明书签署日，发行人实际控制人为甲。'])
            self.make_doc(tar,['情况说明','河南测试有限公司（以下简称“公司”）说明如下：','截至本说明性文件出具日，公司实际控制人为乙。','特此说明。'])
            r=analyze(src,[(str(tar),tar.name)])
            p=next(p for p in r['proposals'] if p['candidates'])
            # Exercise the review-export state machine directly using a real current
            # candidate.  Review is allowed to copy only an explicitly chosen item.
            p['status']='review';p['decision']='keep';p['selected_source_index']=None
            rep=export(r,{p['id']:{'decision':'update'}},out,'柒')
            self.assertTrue(any('必须明确选择当前来源候选' in b['reason'] for b in rep['blocked']))
            out2=td/'out2';si=p['candidates'][0]['source_index']
            rep2=export(r,{p['id']:{'decision':'update','source_index':si}},out2,'柒')
            self.assertFalse(rep2['blocked']);self.assertTrue(rep2['files'])
            rec=next(a for a in rep2['files'][0]['applied_records'] if a['id']==p['id'])
            self.assertTrue(rec['chosen_by_user']);self.assertEqual(rec['source_index'],si)

    def test_source_or_target_change_invalidates_old_analysis(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx';out=td/'out'
            self.make_doc(src,['河南测试有限公司募集说明书','第四节 情况','发行人注册资本为100亿元。'])
            self.make_doc(tar,['情况说明','河南测试有限公司（以下简称“公司”）说明如下：','公司注册资本为90亿元。','特此说明。'])
            r=analyze(src,[(str(tar),tar.name)])
            d=PyDoc(src);d.add_paragraph('来源已变化');d.save(src)
            rep=export(r,{},out,'柒')
            self.assertTrue(any('募集说明书在分析后发生变化' in b['reason'] for b in rep['blocked']))

            # New analysis, then mutate target instead.
            self.make_doc(src,['河南测试有限公司募集说明书','第四节 情况','发行人注册资本为100亿元。'])
            r2=analyze(src,[(str(tar),tar.name)])
            d=PyDoc(tar);d.add_paragraph('目标已变化');d.save(tar)
            rep2=export(r2,{},td/'out2','柒')
            self.assertTrue(any('说明性文件在分析后发生变化' in b['reason'] for b in rep2['blocked']))


@unittest.skipUnless(os.environ.get('EXPLAIN_REAL_PROSPECTUS') and os.environ.get('EXPLAIN_REAL_TARGETS'),'real fixture env not set')
class RealAirportTests(unittest.TestCase):
    def test_real_materials(self):
        with tempfile.TemporaryDirectory() as td:
            targets=unpack_targets(os.environ['EXPLAIN_REAL_TARGETS'],Path(td)/'targets')
            r=analyze(os.environ['EXPLAIN_REAL_PROSPECTUS'],targets)
            self.assertEqual(r['summary']['documents'],29)
            self.assertGreater(r['summary']['source_updates'],100)
            self.assertGreater(r['summary']['wording_updates'],20)
            self.assertLess(r['summary']['review'],50)


class RetentionPolicyTests(unittest.TestCase):
    def make_doc(self,path,paras=(),table=None):
        d=PyDoc()
        for p in paras:d.add_paragraph(p)
        if table:
            t=d.add_table(rows=len(table),cols=len(table[0]))
            for i,row in enumerate(table):
                for j,v in enumerate(row):t.cell(i,j).text=str(v)
        d.save(path)

    def test_partial_candidate_is_completed_retention_not_human_work(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx'
            self.make_doc(src,['河南测试有限公司募集说明书','报告期内，公司与主要客户发生业务往来时不存在严重违约现象。'])
            self.make_doc(tar,['说明','报告期内，公司与主要客户发生业务往来时不存在严重违约情况、不存在债务逾期未偿还的情况。','特此说明。'])
            r=analyze(src,[(str(tar),tar.name)])
            p=next(p for p in r['proposals'] if '债务逾期未偿还' in p['old_text'])
            self.assertEqual(p['status'],'retained')
            self.assertFalse(p['review_required'])
            self.assertEqual(p['candidates'],[])
            self.assertTrue(p.get('correspondence_search',{}).get('candidates'))

    def test_no_candidate_is_completed_retention(self):
        with tempfile.TemporaryDirectory() as td:
            td=Path(td);src=td/'src.docx';tar=td/'tar.docx'
            self.make_doc(src,['河南测试有限公司募集说明书','公司注册资本为100亿元。'])
            self.make_doc(tar,['说明','本事项系说明性文件自身分析结论，与募集说明书没有对应披露。','特此说明。'])
            r=analyze(src,[(str(tar),tar.name)])
            p=next(p for p in r['proposals'] if '自身分析结论' in p['old_text'])
            self.assertEqual(p['status'],'retained')
            self.assertEqual(p['decision'],'keep')

if __name__=='__main__':unittest.main()
